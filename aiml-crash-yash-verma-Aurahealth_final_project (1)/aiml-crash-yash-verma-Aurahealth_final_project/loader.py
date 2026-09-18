import io
import logging
import os
import re
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional

import docx
import pypdf

# Configure logging
logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)


@dataclass
class Document:
    source: str
    text: str
    metadata: Dict[str, Any] = field(default_factory=dict)

    @property
    def patient_id(self) -> Optional[str]:
        return self.metadata.get("patient_id")

    @property
    def year(self) -> Optional[int]:
        return self.metadata.get("year")

    @property
    def document_type(self) -> Optional[str]:
        return self.metadata.get("document_type")


def extract_document_metadata(
    filename: str, text: str, default_patient_id: str = "SYN-PAT-001"
) -> Dict[str, Any]:
    """Extract patient ID, year, and document type from filename and text content."""
    meta: Dict[str, Any] = {
        "source": filename,
        "data_type": "synthetic",
        "document_type": "general_document",
        "patient_id": default_patient_id,
        "year": None,
    }

    # Extract patient ID from filename or text header
    pat_match = re.search(r"(SYN-PAT-\d+)", filename, re.IGNORECASE)
    if pat_match:
        meta["patient_id"] = pat_match.group(1).upper()
    else:
        header_pat = re.search(r"PATIENT\s*ID\s*:\s*([A-Za-z0-9\-]+)", text, re.IGNORECASE)
        if header_pat:
            meta["patient_id"] = header_pat.group(1).strip().upper()

    # Extract 4-digit year from filename (e.g. 2010_SYN-PAT-001.txt)
    year_match = re.search(r"(?:^|[_\-\s])(20\d\d)(?:[_\-\s]|\.|$)", filename)
    if year_match:
        meta["year"] = int(year_match.group(1))
    else:
        header_year = re.search(r"YEAR\s*:\s*(20\d\d)", text, re.IGNORECASE)
        if header_year:
            meta["year"] = int(header_year.group(1))

    if meta["year"] is not None:
        meta["document_type"] = "yearly_medical_history"
    elif "readme" in filename.lower():
        meta["document_type"] = "dataset_documentation"

    return meta


def extract_text_from_bytes(filename: str, file_bytes: bytes) -> str:
    """Extract and normalize text from raw bytes based on file extension (.txt, .pdf, .docx)."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        raw = file_bytes.decode("utf-8", errors="ignore")
        return raw.replace("\r\n", "\n").replace("\r", "\n")

    elif ext == ".pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            raw = "\n\n".join(text_parts)
            return raw.replace("\r\n", "\n").replace("\r", "\n")
        except Exception as e:
            logger.error(f"Failed to extract text from PDF file {filename}: {e}")
            raise ValueError(f"Failed to read PDF file {filename}: {e}")

    elif ext == ".docx":
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            text_parts = [para.text for para in doc.paragraphs if para.text]
            raw = "\n\n".join(text_parts)
            return raw.replace("\r\n", "\n").replace("\r", "\n")
        except Exception as e:
            logger.error(f"Failed to extract text from DOCX file {filename}: {e}")
            raise ValueError(f"Failed to read DOCX file {filename}: {e}")

    else:
        raise ValueError(f"Unsupported file format: '{ext}'. Only .txt, .pdf, and .docx are supported.")


def load_document_from_bytes(
    filename: str, file_bytes: bytes, default_patient_id: str = "SYN-PAT-001"
) -> Document:
    """Create a Document object from uploaded file bytes with extracted metadata."""
    text = extract_text_from_bytes(filename, file_bytes)
    meta = extract_document_metadata(filename, text, default_patient_id=default_patient_id)
    return Document(source=filename, text=text, metadata=meta)


def load_documents(
    data_dir: str = "SYN-PAT-001_15yr_RAG_medical_history",
    default_patient_id: str = "SYN-PAT-001",
) -> List[Document]:
    """Load all supported files (.txt, .pdf, .docx) from data_dir into Document objects with metadata."""
    docs = []
    if not os.path.exists(data_dir):
        logger.warning(f"Data directory '{data_dir}' does not exist.")
        return docs

    for fname in sorted(os.listdir(data_dir)):
        ext = os.path.splitext(fname)[1].lower()
        if ext not in [".txt", ".pdf", ".docx"]:
            continue
        fpath = os.path.join(data_dir, fname)
        try:
            with open(fpath, "rb") as f:
                file_bytes = f.read()
            text = extract_text_from_bytes(fname, file_bytes)
            if text.strip():
                meta = extract_document_metadata(fname, text, default_patient_id=default_patient_id)
                docs.append(Document(source=fname, text=text, metadata=meta))
                logger.info(
                    f"Loaded document: {fname} (Year: {meta.get('year')}, "
                    f"Patient: {meta.get('patient_id')}, Chars: {len(text)})"
                )
            else:
                logger.warning(f"Document {fname} is empty after text extraction.")
        except Exception as e:
            logger.error(f"Error loading document {fname}: {e}")

    return docs


if __name__ == "__main__":
    docs = load_documents("SYN-PAT-001_15yr_RAG_medical_history")
    print(f"Loaded {len(docs)} documents.")
    for d in docs:
        print(f"[{d.year}] {d.patient_id} - {d.source} ({len(d.text)} chars)")
